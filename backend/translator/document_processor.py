"""Document processing for PDF, DOCX, and TXT files."""
import logging
from pathlib import Path
from typing import List, BinaryIO
import io

import PyPDF2
from docx import Document

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Process and translate various document formats."""
    
    SUPPORTED_FORMATS = [".pdf", ".docx", ".txt"]
    MAX_CHUNK_SIZE = 4000  # characters
    
    @staticmethod
    def is_supported(filename: str) -> bool:
        """Check if file format is supported."""
        suffix = Path(filename).suffix.lower()
        return suffix in DocumentProcessor.SUPPORTED_FORMATS
    
    @staticmethod
    def extract_text_from_pdf(file: BinaryIO) -> str:
        """Extract text from PDF file."""
        try:
            pdf_reader = PyPDF2.PdfReader(file)
            text_parts = []
            
            for page_num, page in enumerate(pdf_reader.pages):
                text = page.extract_text()
                if text:
                    text_parts.append(text)
                logger.debug(f"Extracted text from page {page_num + 1}")
            
            full_text = "\n\n".join(text_parts)
            logger.info(f"PDF extraction complete: {len(full_text)} characters")
            return full_text
        except Exception as e:
            logger.error(f"PDF extraction error: {e}")
            raise ValueError(f"Failed to extract text from PDF: {e}")
    
    @staticmethod
    def extract_text_from_docx(file: BinaryIO) -> str:
        """Extract text from DOCX file."""
        try:
            doc = Document(file)
            text_parts = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            full_text = "\n\n".join(text_parts)
            logger.info(f"DOCX extraction complete: {len(full_text)} characters")
            return full_text
        except Exception as e:
            logger.error(f"DOCX extraction error: {e}")
            raise ValueError(f"Failed to extract text from DOCX: {e}")
    
    @staticmethod
    def extract_text_from_txt(file: BinaryIO) -> str:
        """Extract text from TXT file."""
        try:
            content = file.read()
            # Try to decode with utf-8, fallback to latin-1
            try:
                text = content.decode('utf-8')
            except UnicodeDecodeError:
                text = content.decode('latin-1')
            
            logger.info(f"TXT extraction complete: {len(text)} characters")
            return text
        except Exception as e:
            logger.error(f"TXT extraction error: {e}")
            raise ValueError(f"Failed to extract text from TXT: {e}")
    
    @staticmethod
    def extract_text(filename: str, file: BinaryIO) -> str:
        """
        Extract text from supported document formats.
        
        Args:
            filename: Name of the file
            file: File-like object
            
        Returns:
            Extracted text
        """
        suffix = Path(filename).suffix.lower()
        
        if suffix == ".pdf":
            return DocumentProcessor.extract_text_from_pdf(file)
        elif suffix == ".docx":
            return DocumentProcessor.extract_text_from_docx(file)
        elif suffix == ".txt":
            return DocumentProcessor.extract_text_from_txt(file)
        else:
            raise ValueError(f"Unsupported file format: {suffix}")
    
    @staticmethod
    def chunk_text(text: str, chunk_size: int = MAX_CHUNK_SIZE) -> List[str]:
        """
        Split text into chunks for translation.
        
        Args:
            text: Text to split
            chunk_size: Maximum size of each chunk
            
        Returns:
            List of text chunks
        """
        if len(text) <= chunk_size:
            return [text]
        
        chunks = []
        paragraphs = text.split("\n\n")
        current_chunk = ""
        
        for paragraph in paragraphs:
            # If single paragraph is larger than chunk_size, split by sentences
            if len(paragraph) > chunk_size:
                sentences = paragraph.split(". ")
                for sentence in sentences:
                    if len(current_chunk) + len(sentence) + 2 <= chunk_size:
                        current_chunk += sentence + ". "
                    else:
                        if current_chunk:
                            chunks.append(current_chunk.strip())
                        current_chunk = sentence + ". "
            # Add paragraph to current chunk if it fits
            elif len(current_chunk) + len(paragraph) + 2 <= chunk_size:
                current_chunk += paragraph + "\n\n"
            # Start new chunk
            else:
                if current_chunk:
                    chunks.append(current_chunk.strip())
                current_chunk = paragraph + "\n\n"
        
        if current_chunk:
            chunks.append(current_chunk.strip())
        
        logger.info(f"Text chunked into {len(chunks)} parts")
        return chunks
    
    @staticmethod
    def save_translated_document(
        original_filename: str,
        translated_text: str,
        output_format: str = "txt"
    ) -> bytes:
        """
        Save translated text to a file.
        
        Args:
            original_filename: Original file name
            translated_text: Translated text
            output_format: Output format (txt, docx)
            
        Returns:
            File content as bytes
        """
        if output_format == "txt":
            return translated_text.encode('utf-8')
        elif output_format == "docx":
            doc = Document()
            for paragraph in translated_text.split("\n\n"):
                if paragraph.strip():
                    doc.add_paragraph(paragraph)
            
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            return buffer.getvalue()
        else:
            raise ValueError(f"Unsupported output format: {output_format}")


# Global processor instance
processor = DocumentProcessor()
